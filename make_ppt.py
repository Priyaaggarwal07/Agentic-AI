# -----------------------------------------------------------------------------
# make_ppt.py
#
# Generates "MCP_Project_Presentation.pptx" — a slide deck that first explains
# MCP (Model Context Protocol) servers, then introduces THIS project (the AI
# Agricultural Weather Intelligence Platform) and how it maps to the MCP idea.
#
# Run it with the project's venv:
#     backend/venv/Scripts/python.exe make_ppt.py
# -----------------------------------------------------------------------------

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---- Colour palette (matches the app's green theme) ----
GREEN = RGBColor(0x10, 0xB9, 0x81)
TEAL = RGBColor(0x0D, 0x94, 0x88)
DARK = RGBColor(0x0F, 0x17, 0x2A)
SLATE = RGBColor(0x33, 0x41, 0x55)
GRAY = RGBColor(0x47, 0x55, 0x69)
LIGHT = RGBColor(0xF1, 0xF5, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CODEBG = RGBColor(0x1E, 0x29, 0x3B)

prs = Presentation()
prs.slide_width = Inches(13.333)   # 16:9 widescreen
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

SW = prs.slide_width
SH = prs.slide_height


# ---- Small helpers ----------------------------------------------------------

def bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def textbox(slide, text, left, top, width, height, size, color,
            bold=False, align=PP_ALIGN.LEFT, font="Calibri", anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = font
    return tb


def bar(slide, color, height=Inches(1.15)):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def bullets(slide, items, left=Inches(0.9), top=Inches(1.6),
            width=Inches(11.5), height=Inches(5.4)):
    """items: list of (text, level). level 0 = bullet, 1 = sub-bullet."""
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for text, level in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        p.space_after = Pt(10)
        r = p.add_run()
        prefix = "•  " if level == 0 else "–  "
        r.text = prefix + text
        r.font.size = Pt(22 if level == 0 else 19)
        r.font.color.rgb = SLATE if level == 0 else GRAY
        r.font.name = "Calibri"
    return tb


# ---- Slide builders ---------------------------------------------------------

def title_slide(title, subtitle, footer):
    s = prs.slides.add_slide(BLANK)
    bg(s, DARK)
    # green accent stripe
    stripe = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(3.55), SW, Inches(0.08))
    stripe.fill.solid(); stripe.fill.fore_color.rgb = GREEN; stripe.line.fill.background()
    textbox(s, title, Inches(0.9), Inches(2.2), Inches(11.5), Inches(1.3),
            44, WHITE, bold=True)
    textbox(s, subtitle, Inches(0.9), Inches(3.8), Inches(11.5), Inches(1.4),
            22, GREEN)
    textbox(s, footer, Inches(0.9), Inches(6.6), Inches(11.5), Inches(0.6),
            14, LIGHT)
    return s


def section_slide(title):
    s = prs.slides.add_slide(BLANK)
    bg(s, GREEN)
    textbox(s, title, Inches(0.9), Inches(3.0), Inches(11.5), Inches(1.5),
            40, WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return s


def content_slide(title, items):
    s = prs.slides.add_slide(BLANK)
    bg(s, WHITE)
    bar(s, GREEN)
    textbox(s, title, Inches(0.6), Inches(0.18), Inches(12.1), Inches(0.8),
            30, WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    bullets(s, items)
    return s


def code_slide(title, code, caption=None):
    s = prs.slides.add_slide(BLANK)
    bg(s, WHITE)
    bar(s, GREEN)
    textbox(s, title, Inches(0.6), Inches(0.18), Inches(12.1), Inches(0.8),
            30, WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.7),
                             Inches(11.5), Inches(3.9))
    box.fill.solid(); box.fill.fore_color.rgb = CODEBG; box.line.fill.background()
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3); tf.margin_top = Inches(0.25)
    first = True
    for line in code.split("\n"):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        r = p.add_run()
        r.text = line if line else " "
        r.font.name = "Consolas"
        r.font.size = Pt(15)
        r.font.color.rgb = RGBColor(0xA7, 0xF3, 0xD0)
    if caption:
        textbox(s, caption, Inches(0.9), Inches(5.8), Inches(11.5), Inches(1.0),
                18, GRAY)
    return s


# ---- Build the deck ---------------------------------------------------------

title_slide(
    "Model Context Protocol (MCP)",
    "And the AI-Powered Agricultural Weather Intelligence Platform for Nepal",
    "A rule-based crop engine + LangGraph agent over real weather & soil APIs",
)

# ===== PART 1: MCP =====
section_slide("Part 1  —  Understanding MCP Servers")

content_slide("What is MCP?", [
    ("Model Context Protocol — an open standard introduced by Anthropic (2024).", 0),
    ("A universal way to connect AI models (LLMs) to external tools, data, and services.", 0),
    ("Often called \"USB-C for AI\": one standard interface instead of custom, one-off integrations.", 0),
    ("Lets an LLM safely call real functions and read real data — beyond its training knowledge.", 0),
])

content_slide("The Problem MCP Solves", [
    ("Before MCP, every (model + tool) integration was custom and non-reusable.", 0),
    ("N models × M tools  =  N×M bespoke connectors to build and maintain.", 0),
    ("Hard to reuse, hard to secure, hard to scale.", 0),
    ("MCP: build a tool once as a \"server\" — any MCP-aware model can use it.", 0),
])

content_slide("MCP Architecture", [
    ("Host — the AI application (e.g. Claude Desktop, an agent, an IDE).", 0),
    ("Client — lives inside the host and speaks the MCP protocol.", 0),
    ("Server — exposes your capabilities (tools, data) to the model.", 0),
    ("Transport — JSON-RPC 2.0 messages over stdio or HTTP.", 0),
    ("Host ⇄ Client ⇄ Server  —  a clean, standardized boundary.", 1),
])

content_slide("What an MCP Server Exposes", [
    ("Tools — functions the model can call (e.g. get_weather, query_db).", 0),
    ("Resources — data the model can read (files, records, documents).", 0),
    ("Prompts — reusable prompt templates the host can offer the user.", 0),
    ("Each capability carries a description so the model knows WHEN to use it.", 0),
])

content_slide("How a Tool Call Works", [
    ("1. The model decides it needs a tool to answer.", 0),
    ("2. The client sends a JSON-RPC \"tools/call\" request to the server.", 0),
    ("3. The server runs the function and returns the result.", 0),
    ("4. The model reads the result and continues its answer.", 0),
    ("Same request/response idea as a REST API — but standardized for LLMs.", 1),
])

# ===== PART 2: The project =====
section_slide("Part 2  —  Introducing My Project")

content_slide("Project Overview", [
    ("AI-Powered Agricultural Weather Intelligence Platform — built for Nepal.", 0),
    ("Click a location on the map → get weather, soil, crop suitability, fertilizer & forecast.", 0),
    ("Stack: React frontend + FastAPI backend + real external APIs + an AI agent.", 0),
    ("Built with LangGraph tools — the same \"agent + tools\" pattern MCP standardizes.", 0),
])

content_slide("System Architecture", [
    ("Frontend — React + Leaflet map; user clicks a coordinate.", 0),
    ("Backend — FastAPI endpoints: /weather /soil /forecast /crop-suitability /ask.", 0),
    ("External data — OpenWeather API and the Nepal NARC Soil API.", 0),
    ("AI Agent — LangGraph + OpenAI decides which backend tools to call.", 0),
])

content_slide("My \"Tools\" (MCP-style)", [
    ("weather_lookup — current weather (OpenWeather).", 0),
    ("soil_lookup — soil texture, pH, and N/P/K levels (Nepal soil API).", 0),
    ("forecast_lookup — a simple 3-day forecast.", 0),
    ("crop_recommendation — rule-based engine scoring 20+ Nepali crops.", 0),
    ("fertilizer_lookup — DAP / Urea / MOP amounts for a crop.", 0),
    ("Each has a clear description so the agent knows when to call it — exactly like MCP tools.", 1),
])

content_slide("The Agentic Layer", [
    ("User asks a plain-English question: \"What should I grow here?\"", 0),
    ("A LangGraph ReAct agent (OpenAI gpt-4o-mini) reads the question.", 0),
    ("It chooses tools, calls them, reads the results, and writes the answer.", 0),
    ("Exposed through a single endpoint: POST /ask.", 0),
])

code_slide(
    "Code: A Tool Definition",
    '@tool\n'
    'def weather_lookup(latitude: float, longitude: float) -> dict:\n'
    '    """Get the CURRENT weather (temperature, humidity, wind,\n'
    '    and city) for a latitude/longitude in Nepal."""\n'
    '    return get_weather(latitude, longitude)',
    caption="The docstring is the tool's description the LLM reads — the same role a description plays in MCP.",
)

content_slide("Example: A Full Question", [
    ("\"What should I grow at 28.29, 83.98 and how much urea?\"", 0),
    ("Agent → crop_recommendation(28.29, 83.98)", 1),
    ("...which fetches weather + soil in parallel (two threads)", 1),
    ("Agent → fertilizer_lookup(28.29, 83.98, \"Rice\")", 1),
    ("Answer: \"Rice (98/100). Apply DAP 2.17, Urea 6.41, MOP 1.67 kg/ropani.\"", 0),
])

content_slide("Performance: Parallel Tool Calls", [
    ("Weather and soil are independent network calls (mostly waiting on I/O).", 0),
    ("Run them on two threads with a ThreadPoolExecutor → the waits overlap.", 0),
    ("Roughly halves the time vs. calling one after the other.", 0),
    ("Python releases the GIL during network I/O, so threads genuinely help here.", 1),
])

content_slide("Tech Stack", [
    ("Frontend — React (Vite), Leaflet, Tailwind CSS, Axios.", 0),
    ("Backend — FastAPI, Python, Requests, python-dotenv.", 0),
    ("AI — LangGraph, LangChain, OpenAI (gpt-4o-mini).", 0),
    ("Data — OpenWeather API, Nepal NARC Soil API.", 0),
])

content_slide("How This Maps to MCP  —  Next Step", [
    ("My LangChain tools ≈ MCP tools: described, callable functions.", 0),
    ("Next step: wrap the same functions in an MCP server (JSON-RPC transport).", 0),
    ("Then ANY MCP-aware client (Claude Desktop, other agents) could reuse them.", 0),
    ("No rewrite of the domain logic — only the transport layer changes.", 1),
])

content_slide("Summary", [
    ("MCP = a standard way to give LLMs tools and data.", 0),
    ("My project already uses the agent + tools pattern that MCP formalizes.", 0),
    ("Real APIs, a rule-based crop engine, and an AI agent for Nepali farmers.", 0),
    ("A clear, low-effort path to exposing it as a true MCP server.", 0),
])

# closing
closing = prs.slides.add_slide(BLANK)
bg(closing, DARK)
textbox(closing, "Thank You", Inches(0.9), Inches(2.7), Inches(11.5), Inches(1.3),
        48, WHITE, bold=True, align=PP_ALIGN.CENTER)
textbox(closing, "Questions?", Inches(0.9), Inches(4.0), Inches(11.5), Inches(0.9),
        26, GREEN, align=PP_ALIGN.CENTER)

OUT = "MCP_Project_Presentation.pptx"
prs.save(OUT)
print(f"Saved {OUT} with {len(prs.slides._sldIdLst)} slides.")
